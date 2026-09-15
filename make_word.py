# make_word.py
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os

def create_resume_docx(input_txt, output_docx):
    try:
        doc = Document()
        for section in doc.sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        title = doc.add_heading('My Optimized Resume', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if not os.path.exists(input_txt):
            print(f"[ERROR] Input file not found: {input_txt}")
            return

        with open(input_txt, 'r', encoding='utf-8') as f:
            content = f.read()

        lines = content.split('\n')
        for line in lines:
            line = line.strip()
            if not line:
                continue
            if line.startswith('##'):
                doc.add_heading(line.replace('##', '').strip(), level=1)
            elif line.startswith('###'):
                doc.add_heading(line.replace('###', '').strip(), level=2)
            else:
                p = doc.add_paragraph(line)
                p.paragraph_format.space_after = Pt(6)
                for run in p.runs:
                    run.font.name = 'Microsoft YaHei'
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
                    run.font.size = Pt(11)

        doc.save(output_docx)
        print(f"[OK] Word resume generated: {output_docx}")

    except Exception as e:
        print(f"[ERROR] {str(e)}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    create_resume_docx("optimized_resume.txt", "resume_final.docx")